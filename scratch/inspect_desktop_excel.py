import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Admin\Desktop\OUTSTANDING RECEIVABLE REPORT WITH POLICY.xlsx', read_only=True)
print('Sheets:', wb.sheetnames)
sheet = wb.active
print('Active sheet:', sheet.title)

for i, row in enumerate(sheet.iter_rows(values_only=True)):
    if i < 5:
        print(f"Row {i+1}: {row[:10]}")
    # Search for SAN
    row_str = ' '.join(str(c or '') for c in row)
    if 'SAN' in row_str and ('ANNAPURNA' in row_str or 'KALPESH' in row_str):
        print(f"SAN Row {i+1}: {row}")
    if i > 6000:
        break
